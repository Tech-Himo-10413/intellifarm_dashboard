from pptx import Presentation

prs = Presentation('IntelliFarm_WASSAN_Presentation.pptx')
for i, slide in enumerate(prs.slides):
    print(f"\n--- Slide {i+1} ---")
    for shape in slide.shapes:
        if shape.has_text_frame:
            text = shape.text.replace('\n', ' | ')
            print(text.encode('utf-8', 'ignore').decode('utf-8'))

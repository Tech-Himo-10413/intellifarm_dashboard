from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

prs = Presentation('IntelliFarm_WASSAN_Presentation.pptx')
width = prs.slide_width
height = prs.slide_height
print(f"Width: {width.inches}, Height: {height.inches}")

# Let's insert new slides before the last slide (index 5)
last_slide = prs.slides[5]

# Actually, python-pptx doesn't have an easy way to insert slides at a specific index. 
# It always appends to the end. But we can move the slides XML around, or we can just 
# append them to the end, then move the last slide to the end.

# Here is how to move a slide to the end:
def move_slide(prs, old_index, new_index):
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    xml_slides.remove(slides[old_index])
    xml_slides.insert(new_index, slides[old_index])

blank_slide_layout = prs.slide_layouts[5] # Usually 5 or 6 is Title Only or Blank. Let's find one.
# We can just use layout 5 (Title only)

images = [
    {"path": "pictures/initial_page.png", "title": "Phase 1: Upload & Health Check"},
    {"path": "pictures/clean_validate.png", "title": "Phase 2: Clean & Validate"},
    {"path": "pictures/dashboard_gen_ai_suggestions.png", "title": "Phase 3: AI Dashboard Suggestions"},
    {"path": "pictures/ai_dashboarrd_generations.png", "title": "Phase 3: AI Dashboard Generations"}
]

for img_info in images:
    slide = prs.slides.add_slide(prs.slide_layouts[5]) # Title only layout
    title_shape = slide.shapes.title
    title_shape.text = img_info["title"]
    
    # Optional styling for title
    for paragraph in title_shape.text_frame.paragraphs:
        paragraph.font.size = Pt(32)
        paragraph.font.color.rgb = RGBColor(0, 102, 102)
        paragraph.alignment = PP_ALIGN.LEFT
    
    # Calculate image dimensions to fit well
    img_path = img_info["path"]
    
    # margins
    left = Inches(0.5)
    top = Inches(1.5)
    max_width = width - Inches(1)
    max_height = height - Inches(2)
    
    # Insert image, let python-pptx scale it proportionally to fit within max width/height if we just specify width or height.
    # We will just pass width and let height scale. Or pass height and let width scale.
    # A safe way is to add the picture with a fixed height and see if width exceeds.
    try:
        pic = slide.shapes.add_picture(img_path, left, top)
        # scale logic
        ratio = pic.width / pic.height
        max_ratio = max_width / max_height
        
        if ratio > max_ratio:
            # wider than max area
            pic.width = max_width
            pic.height = int(max_width / ratio)
        else:
            # taller than max area
            pic.height = max_height
            pic.width = int(max_height * ratio)
            
        # center it horizontally
        pic.left = int((width - pic.width) / 2)
        
    except Exception as e:
        print(f"Error adding {img_path}: {e}")

# Now we added 4 slides at the end (indices 6, 7, 8, 9).
# The original conclusion slide is at index 5.
# We want the conclusion slide to be at the very end.
move_slide(prs, 5, 9)

prs.save('IntelliFarm_WASSAN_Presentation.pptx')
print("Presentation updated successfully.")

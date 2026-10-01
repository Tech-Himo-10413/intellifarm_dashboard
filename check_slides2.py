from pptx import Presentation

prs = Presentation('IntelliFarm_WASSAN_Presentation.pptx')
for i, slide in enumerate(prs.slides):
    title = slide.shapes.title.text if slide.shapes.title and slide.shapes.title.has_text_frame else "No Title"
    print(f"Slide {i+1}: {title.replace(chr(10), ' ')}")
